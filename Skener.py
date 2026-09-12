from __future__ import annotations
import sys
import os
import tempfile
from typing import Optional, Callable

from PyQt6.QtWidgets import (
    QApplication, QWidget, QVBoxLayout, QHBoxLayout, QPushButton, QLabel,
    QComboBox, QMessageBox, QFileDialog, QDialog, QListWidget,
    QProgressDialog, QTextEdit, QCheckBox, QAbstractItemView, QGroupBox,
    QInputDialog, QScrollArea, QRadioButton, QDialogButtonBox, QButtonGroup
)
from PyQt6.QtGui import QPixmap, QShortcut, QKeySequence
from PyQt6.QtCore import Qt, QEventLoop, QSettings, QTimer, QStandardPaths
import re

# QAccessible není v PyQt6 6.11+ exponován v Python API (ověřeno
# ImportError: cannot import name 'QAccessible' from 'PyQt6.QtGui').
# Qt interně posílá Focus event automaticky při setFocus(), takže
# explicitní QAccessible.updateAccessibility() není nutné. Pokud by
# budoucí verze PyQt6 QAccessible zpřístupnila, lze jej volat
# podmíněně přes importlib – viz ScanApp._set_initial_focus().
from PIL import Image, ImageQt, ImageFilter, ImageOps
from docx import Document
import fitz
from accessible_output2.outputs.auto import Auto

from scanner_engine import NAPS2Scanner, ScanThread
from ocr_engine import create_ocr_thread, OcrResult
from macro import Macro, MacroManager, PipelineRunner
from macro_editor import MacroEditorDialog

# ------------------ Hlasový výstup ------------------
_speaker = Auto()

def speak(text: str) -> None:
    _speaker.output(text)

# ------------------ Předzpracování obrazu ------------------
def preprocess_image(img: Image.Image) -> Image.Image:
    if img.mode != "L":
        img = img.convert("L")
    img = img.filter(ImageFilter.SHARPEN)
    img = ImageOps.autocontrast(img, cutoff=2)
    return img


def alt_preprocess_image(img: Image.Image) -> Image.Image:
    if img.mode != "L":
        img = img.convert("L")
    img = img.filter(ImageFilter.MedianFilter(3))
    img = ImageOps.autocontrast(img, cutoff=1)
    return img

# ------------------ Náhled dialog ------------------
class PreviewDialog(QDialog):
    def __init__(self, pil_img: Image.Image, page_num: int = 1) -> None:
        super().__init__()
        self.setWindowTitle("Náhled skenu")
        self.image = pil_img
        self.page_num = page_num

        self.label = QLabel()
        self.update_image()

        layout = QVBoxLayout()
        layout.addWidget(self.label)

        btn_layout = QHBoxLayout()
        btn_rotate = QPushButton("Otočit")
        btn_rotate.setAccessibleName("Otočit obrázek")
        btn_rotate.clicked.connect(self.rotate_image)
        btn_layout.addWidget(btn_rotate)

        btn_ok = QPushButton("Potvrdit náhled")
        btn_ok.setAccessibleName("Potvrdit náhled")
        btn_ok.clicked.connect(self.accept)
        btn_layout.addWidget(btn_ok)

        layout.addLayout(btn_layout)
        self.setLayout(layout)

        QTimer.singleShot(300, lambda: speak(
            f"Náhled stránky {page_num}. Otočte obrázek nebo potvrďte náhled."
        ))

    def update_image(self) -> None:
        qimg = ImageQt.ImageQt(self.image)
        pix = QPixmap.fromImage(qimg).scaled(
            500, 700, Qt.AspectRatioMode.KeepAspectRatio
        )
        self.label.setPixmap(pix)

    def rotate_image(self) -> None:
        self.image = self.image.rotate(-90, expand=True)
        self.update_image()
        speak("Obrázek otočen.")

# ------------------ Dialog pro editaci OCR výsledku ------------------
class OcrPreviewDialog(QDialog):
    def __init__(self, text: str, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setWindowTitle("Náhled OCR textu")
        self.setMinimumSize(600, 400)

        self.text_edit = QTextEdit()
        self.text_edit.setPlainText(text)
        self.text_edit.setAccessibleName("OCR text k editaci")

        btn_layout = QHBoxLayout()
        btn_read = QPushButton("Přečíst text")
        btn_read.setAccessibleName("Přečíst text hlasem")
        btn_read.clicked.connect(self.read_text)
        btn_layout.addWidget(btn_read)

        btn_save = QPushButton("Uložit")
        btn_save.setAccessibleName("Uložit a pokračovat")
        btn_save.setDefault(True)
        btn_save.clicked.connect(self.accept)
        btn_layout.addWidget(btn_save)

        btn_cancel = QPushButton("Zrušit")
        btn_cancel.clicked.connect(self.reject)
        btn_layout.addWidget(btn_cancel)

        layout = QVBoxLayout()
        layout.addWidget(self.text_edit)
        layout.addLayout(btn_layout)
        self.setLayout(layout)

    def read_text(self) -> None:
        speak(self.text_edit.toPlainText())

    def showEvent(self, event) -> None:
        super().showEvent(event)
        self.text_edit.setFocus()

    def get_text(self) -> str:
        return self.text_edit.toPlainText()


# ------------------ Pomocné funkce pro stránkování ------------------
_PAGE_HEADER_RE = re.compile(r"^---\s*Stránka\s+(\d+)\s*---\s*$", re.MULTILINE)


def _split_text_by_page_headers(text: str) -> list[str] | None:
    """Rozdělí text podle headerů '--- Stránka N ---' na list per-page.

    Vrátí None pokud text neobsahuje žádné headery nebo je neplatný.
    """
    if not text or "--- Stránka" not in text:
        return None
    # Najdi všechny headery
    matches = list(_PAGE_HEADER_RE.finditer(text))
    if not matches:
        return None
    pages: list[str] = []
    for idx, m in enumerate(matches):
        start = m.end()
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
        page_content = text[start:end].strip("\n")
        # Odstraň přebytečné okolní nové řádky, ale zachovej vnitřní
        # strip jen koncové \n, pak strip surrounding whitespace per page
        pages.append(page_content.strip())
    return pages


# ------------------ Dialog pro výběr formátu uložení (legacy) ------------------
class SaveDocumentDialog(QDialog):
    """Přístupný dialog 'Uložit dokument' – výběr formátu s lidskými názvy (ponechán pro kompatibilitu makro)."""

    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setWindowTitle("Uložit dokument")
        self.setMinimumWidth(520)
        self._selected_format: str = "txt"

        layout = QVBoxLayout(self)

        lbl_info = QLabel("Vyberte formát, ve kterém chcete dokument uložit.")
        lbl_info.setWordWrap(True)
        lbl_info.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        layout.addWidget(lbl_info)

        # Radio buttons s lidskými názvy a popisem
        self.rb_txt = QRadioButton("Textový dokument (.txt)")
        self.rb_txt.setAccessibleName("Textový dokument (.txt)")
        self.rb_txt.setAccessibleDescription("Pouhý text bez zachování vzhledu stránky.")
        self.rb_txt.setChecked(True)
        layout.addWidget(self.rb_txt)
        lbl_txt_desc = QLabel("Pouhý text bez zachování vzhledu stránky.")
        lbl_txt_desc.setWordWrap(True)
        lbl_txt_desc.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        lbl_txt_desc.setStyleSheet("color: palette(mid); margin-left: 22px;")
        layout.addWidget(lbl_txt_desc)

        self.rb_docx = QRadioButton("Word dokument (.docx)")
        self.rb_docx.setAccessibleName("Word dokument (.docx)")
        self.rb_docx.setAccessibleDescription("Upravitelný dokument pro Microsoft Word a další kompatibilní programy.")
        layout.addWidget(self.rb_docx)
        lbl_docx_desc = QLabel("Upravitelný dokument pro Microsoft Word a další kompatibilní programy.")
        lbl_docx_desc.setWordWrap(True)
        lbl_docx_desc.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        lbl_docx_desc.setStyleSheet("color: palette(mid); margin-left: 22px;")
        layout.addWidget(lbl_docx_desc)

        self.rb_pdf = QRadioButton("PDF s OCR (.pdf)")
        self.rb_pdf.setAccessibleName("PDF s OCR (.pdf)")
        self.rb_pdf.setAccessibleDescription("Dokument se zachovaným vzhledem stránek a skrytou textovou vrstvou pro vyhledávání a odečítání textu.")
        layout.addWidget(self.rb_pdf)
        lbl_pdf_desc = QLabel("Dokument se zachovaným vzhledem stránek a skrytou textovou vrstvou pro vyhledávání a odečítání textu.")
        lbl_pdf_desc.setWordWrap(True)
        lbl_pdf_desc.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        lbl_pdf_desc.setStyleSheet("color: palette(mid); margin-left: 22px;")
        layout.addWidget(lbl_pdf_desc)

        # Button group pro logické propojení
        self._group = QButtonGroup(self)
        self._group.addButton(self.rb_txt, 0)
        self._group.addButton(self.rb_docx, 1)
        self._group.addButton(self.rb_pdf, 2)

        btn_box = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel, self)
        btn_ok = btn_box.button(QDialogButtonBox.StandardButton.Ok)
        btn_ok.setText("Uložit")
        btn_ok.setAccessibleName("Uložit v zvoleném formátu")
        btn_ok.setDefault(True)
        btn_cancel = btn_box.button(QDialogButtonBox.StandardButton.Cancel)
        btn_cancel.setText("Zrušit")
        btn_cancel.setAccessibleName("Zrušit ukládání")
        btn_box.accepted.connect(self.accept)
        btn_box.rejected.connect(self.reject)
        layout.addWidget(btn_box)

        # Tab order je přirozený: rb_txt -> rb_docx -> rb_pdf -> Ok -> Cancel
        self.setTabOrder(self.rb_txt, self.rb_docx)
        self.setTabOrder(self.rb_docx, self.rb_pdf)
        self.setTabOrder(self.rb_pdf, btn_ok)
        self.setTabOrder(btn_ok, btn_cancel)

    def showEvent(self, event) -> None:
        super().showEvent(event)
        # Počáteční fokus na první volbu
        self.rb_txt.setFocus(Qt.FocusReason.OtherFocusReason)
        speak("Dialog Uložit dokument. Vyberte formát, ve kterém chcete dokument uložit.")

    def selected_format(self) -> str:
        if self.rb_docx.isChecked():
            return "docx"
        if self.rb_pdf.isChecked():
            return "pdf"
        return "txt"


# ------------------ Dvoustupňové dialogy pro oddělené ukládání ------------------
class ChooseSaveKindDialog(QDialog):
    """První krok 'Co chcete uložit?' – odděluje uložení bez OCR a s OCR."""

    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setWindowTitle("Co chcete uložit?")
        self.setMinimumWidth(520)
        layout = QVBoxLayout(self)
        lbl_info = QLabel("Co chcete uložit?")
        lbl_info.setWordWrap(True)
        lbl_info.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        # accessible name for NVDA heading
        lbl_info.setAccessibleName("Co chcete uložit")
        layout.addWidget(lbl_info)

        self.rb_images = QRadioButton("Naskenované stránky")
        self.rb_images.setAccessibleName("Naskenované stránky")
        self.rb_images.setAccessibleDescription("Uloží stránky bez OCR jako PDF. Uloží stránky jako obrázky. OCR nebude proveden.")
        self.rb_images.setChecked(True)
        layout.addWidget(self.rb_images)
        lbl_images_desc = QLabel("Uloží stránky bez OCR jako PDF.")
        lbl_images_desc.setWordWrap(True)
        lbl_images_desc.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        lbl_images_desc.setStyleSheet("color: palette(mid); margin-left: 22px;")
        layout.addWidget(lbl_images_desc)
        lbl_images_detail = QLabel("Uloží stránky jako obrázky. OCR nebude proveden.")
        lbl_images_detail.setWordWrap(True)
        lbl_images_detail.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        lbl_images_detail.setStyleSheet("color: palette(mid); margin-left: 22px; font-style: italic;")
        layout.addWidget(lbl_images_detail)

        self.rb_ocr = QRadioButton("OCR dokument")
        self.rb_ocr.setAccessibleName("OCR dokument")
        self.rb_ocr.setAccessibleDescription("Uloží již rozpoznaný text jako TXT, DOCX nebo PDF s OCR. Vyžaduje předchozí OCR.")
        layout.addWidget(self.rb_ocr)
        lbl_ocr_desc = QLabel("Uloží již rozpoznaný text jako TXT, DOCX nebo PDF s OCR.")
        lbl_ocr_desc.setWordWrap(True)
        lbl_ocr_desc.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        lbl_ocr_desc.setStyleSheet("color: palette(mid); margin-left: 22px;")
        layout.addWidget(lbl_ocr_desc)

        self._group = QButtonGroup(self)
        self._group.addButton(self.rb_images, 0)
        self._group.addButton(self.rb_ocr, 1)

        btn_box = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel, self)
        btn_ok = btn_box.button(QDialogButtonBox.StandardButton.Ok)
        btn_ok.setText("Pokračovat")
        btn_ok.setAccessibleName("Pokračovat v ukládání")
        btn_ok.setDefault(True)
        btn_cancel = btn_box.button(QDialogButtonBox.StandardButton.Cancel)
        btn_cancel.setText("Zrušit")
        btn_cancel.setAccessibleName("Zrušit ukládání")
        btn_box.accepted.connect(self.accept)
        btn_box.rejected.connect(self.reject)
        layout.addWidget(btn_box)
        self.setTabOrder(self.rb_images, self.rb_ocr)
        self.setTabOrder(self.rb_ocr, btn_ok)
        self.setTabOrder(btn_ok, btn_cancel)
        self._btn_ok = btn_ok
        self._btn_cancel = btn_cancel

    def showEvent(self, event) -> None:
        super().showEvent(event)
        self.rb_images.setFocus(Qt.FocusReason.OtherFocusReason)
        speak("Dialog Co chcete uložit. Vyberte Naskenované stránky bez OCR nebo OCR dokument s rozpoznaným textem.")

    def selected_kind(self) -> str:
        if self.rb_ocr.isChecked():
            return "ocr"
        return "images"


class ChooseOcrFormatDialog(QDialog):
    """Druhý krok – výběr konkrétního OCR formátu."""

    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setWindowTitle("Uložit OCR dokument")
        self.setMinimumWidth(520)
        layout = QVBoxLayout(self)
        lbl_info = QLabel("Vyberte formát OCR dokumentu.")
        lbl_info.setWordWrap(True)
        lbl_info.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        layout.addWidget(lbl_info)

        self.rb_txt = QRadioButton("Textový dokument (.txt)")
        self.rb_txt.setAccessibleName("Textový dokument (.txt)")
        self.rb_txt.setAccessibleDescription("Pouhý text bez zachování vzhledu stránky. Vyžaduje OCR.")
        self.rb_txt.setChecked(True)
        layout.addWidget(self.rb_txt)
        lbl_txt_desc = QLabel("Pouhý text bez zachování vzhledu stránky.")
        lbl_txt_desc.setWordWrap(True)
        lbl_txt_desc.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        lbl_txt_desc.setStyleSheet("color: palette(mid); margin-left: 22px;")
        layout.addWidget(lbl_txt_desc)

        self.rb_docx = QRadioButton("Word dokument (.docx)")
        self.rb_docx.setAccessibleName("Word dokument (.docx)")
        self.rb_docx.setAccessibleDescription("Upravitelný dokument pro Microsoft Word. Vyžaduje OCR.")
        layout.addWidget(self.rb_docx)
        lbl_docx_desc = QLabel("Upravitelný dokument pro Microsoft Word a další kompatibilní programy.")
        lbl_docx_desc.setWordWrap(True)
        lbl_docx_desc.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        lbl_docx_desc.setStyleSheet("color: palette(mid); margin-left: 22px;")
        layout.addWidget(lbl_docx_desc)

        self.rb_pdf_ocr = QRadioButton("PDF s rozpoznaným textem (.pdf)")
        self.rb_pdf_ocr.setAccessibleName("PDF s rozpoznaným textem (.pdf)")
        self.rb_pdf_ocr.setAccessibleDescription("Zachová vzhled naskenovaných stránek a přidá textovou vrstvu z OCR. Vyžaduje OCR.")
        layout.addWidget(self.rb_pdf_ocr)
        lbl_pdf_desc = QLabel("Zachová vzhled naskenovaných stránek a přidá textovou vrstvu z OCR.")
        lbl_pdf_desc.setWordWrap(True)
        lbl_pdf_desc.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        lbl_pdf_desc.setStyleSheet("color: palette(mid); margin-left: 22px;")
        layout.addWidget(lbl_pdf_desc)

        self._group = QButtonGroup(self)
        self._group.addButton(self.rb_txt, 0)
        self._group.addButton(self.rb_docx, 1)
        self._group.addButton(self.rb_pdf_ocr, 2)

        btn_box = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel, self)
        btn_ok = btn_box.button(QDialogButtonBox.StandardButton.Ok)
        btn_ok.setText("Uložit")
        btn_ok.setAccessibleName("Uložit v zvoleném OCR formátu")
        btn_ok.setDefault(True)
        btn_cancel = btn_box.button(QDialogButtonBox.StandardButton.Cancel)
        btn_cancel.setText("Zrušit")
        btn_cancel.setAccessibleName("Zrušit ukládání")
        btn_box.accepted.connect(self.accept)
        btn_box.rejected.connect(self.reject)
        layout.addWidget(btn_box)
        self.setTabOrder(self.rb_txt, self.rb_docx)
        self.setTabOrder(self.rb_docx, self.rb_pdf_ocr)
        self.setTabOrder(self.rb_pdf_ocr, btn_ok)
        self.setTabOrder(btn_ok, btn_cancel)
        self._btn_ok = btn_ok
        self._btn_cancel = btn_cancel

    def showEvent(self, event) -> None:
        super().showEvent(event)
        self.rb_txt.setFocus(Qt.FocusReason.OtherFocusReason)
        speak("Dialog Uložit OCR dokument. Vyberte formát TXT, DOCX nebo PDF s rozpoznaným textem.")

    def selected_format(self) -> str:
        if self.rb_docx.isChecked():
            return "docx"
        if self.rb_pdf_ocr.isChecked():
            return "pdf_ocr"
        return "txt"


# ------------------ Hlavní aplikace ------------------
class ScanApp(QWidget):
    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("Hlasový skener – NAPS2 + OCR")
        self.setMinimumSize(500, 600)

        self.settings = QSettings("HlasovySkener", "HlasovySkener")
        self.scanner = NAPS2Scanner()
        self.scanned_images: list[Image.Image] = []
        self.raw_scanned_images: list[Image.Image] = []
        self.last_ocr_results: list[list[OcrResult]] = []
        self.last_scanned_image: Optional[bool] = None
        self._scan_cancelled = False
        self._ocr_mode = "interactive"
        self._last_text = ""
        self._last_original_text = ""
        self._last_edited_pages: list[str] | None = None
        self._diacritics_failed = False
        self._diacritics_enabled_cache = False
        self._ocr_total_pages: int = 0
        self._ocr_last_announced_page: int = 0
        # Pro dvoustupňové ukládání – zachování volby přes OCR
        self._pending_save_format: str | None = None

        self.macro_manager = MacroManager()
        self.macro_manager.load_all()
        self._macro_runner = PipelineRunner(self)
        self._macro_shortcuts: list[QShortcut] = []
        # Flag pro jednorázové nastavení počátečního fokusu po zobrazení okna
        self._initial_focus_done: bool = False

        self._build_ui()
        self._setup_shortcuts()
        self._load_settings()

    def speak(self, text: str) -> None:
        speak(text)

    # ---------- UI ----------
    def _build_ui(self) -> None:
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QScrollArea.Shape.NoFrame)
        # QScrollArea nesmí krást počáteční fokus – je to kontejner
        scroll.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        # Viewport scroll area také nesmí být focusable (Qt default může být StrongFocus u některých stylů)
        scroll.viewport().setFocusPolicy(Qt.FocusPolicy.NoFocus)

        container = QWidget()
        container.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        layout = QVBoxLayout(container)
        layout.setSpacing(12)
        layout.setContentsMargins(15, 15, 15, 15)

        # -- Nastavení OCR --
        ocr_group = QGroupBox("Nastavení OCR")
        ocr_layout = QVBoxLayout()

        lbl_lang = QLabel("Jazyk OCR:")
        self.lang_combo = QComboBox()
        self.lang_combo.setAccessibleName("Jazyk OCR")
        self.lang_combo.addItems([
            "ces (Čeština)", "eng (Angličtina)", "deu (Němčina)",
            "fra (Francouzština)", "ita (Italština)", "pol (Polština)"
        ])
        lbl_lang.setBuddy(self.lang_combo)
        ocr_layout.addWidget(lbl_lang)
        ocr_layout.addWidget(self.lang_combo)

        lbl_engine = QLabel("OCR engine:")
        self.engine_combo = QComboBox()
        self.engine_combo.addItems(["Tesseract", "EasyOCR"])
        self.engine_combo.setAccessibleName("OCR engine")
        lbl_engine.setBuddy(self.engine_combo)
        ocr_layout.addWidget(lbl_engine)
        ocr_layout.addWidget(self.engine_combo)

        self.preprocess_cb = QCheckBox("Předzpracovat obraz (doostřit, kontrast)")
        self.preprocess_cb.setAccessibleName("Předzpracování obrazu")
        self.preprocess_cb.setChecked(True)
        ocr_layout.addWidget(self.preprocess_cb)

        lbl_diac = QLabel("Oprava české diakritiky:")
        self.diacritics_cb = QCheckBox("Oprava české diakritiky")
        self.diacritics_cb.setAccessibleName("Oprava české diakritiky")
        self.diacritics_cb.setChecked(False)
        self.diacritics_cb.toggled.connect(self._on_diacritics_toggled)
        lbl_diac.setBuddy(self.diacritics_cb)
        ocr_layout.addWidget(lbl_diac)
        ocr_layout.addWidget(self.diacritics_cb)

        ocr_group.setLayout(ocr_layout)
        layout.addWidget(ocr_group)

        # -- Nastavení skeneru --
        scan_group = QGroupBox("Nastavení skeneru")
        scan_layout = QVBoxLayout()

        self.batch_cb = QCheckBox("Dávkový režim (automaticky všechny stránky)")
        self.batch_cb.setAccessibleName("Dávkový režim")
        scan_layout.addWidget(self.batch_cb)

        lbl_device = QLabel("Vyber profil NAPS2:")
        self.device_combo = QComboBox()
        self.device_combo.setAccessibleName("Vyber profil NAPS2")
        lbl_device.setBuddy(self.device_combo)
        self.reload_devices()
        scan_layout.addWidget(lbl_device)
        scan_layout.addWidget(self.device_combo)

        lbl_source = QLabel("Zdroj papíru:")
        self.source_combo = QComboBox()
        self.source_combo.setAccessibleName("Zdroj papíru")
        self.source_combo.addItems(["Sklo", "Podavač"])
        lbl_source.setBuddy(self.source_combo)
        scan_layout.addWidget(lbl_source)
        scan_layout.addWidget(self.source_combo)

        lbl_dpi = QLabel("Rozlišení DPI:")
        self.dpi_combo = QComboBox()
        self.dpi_combo.setAccessibleName("Rozlišení DPI")
        self.dpi_combo.setEditable(False)
        self.dpi_combo.addItem("200 DPI – rychlé skenování", 200)
        self.dpi_combo.addItem("300 DPI – doporučeno pro OCR", 300)
        self.dpi_combo.addItem("400 DPI – menší text", 400)
        self.dpi_combo.addItem("600 DPI – velmi malý text", 600)
        self.dpi_combo.addItem("1200 DPI – vysoká kvalita", 1200)
        idx = self.dpi_combo.findData(300)
        if idx >= 0:
            self.dpi_combo.setCurrentIndex(idx)
        lbl_dpi.setBuddy(self.dpi_combo)
        scan_layout.addWidget(lbl_dpi)
        scan_layout.addWidget(self.dpi_combo)

        lbl_color = QLabel("Režim:")
        self.color_combo = QComboBox()
        self.color_combo.setAccessibleName("Režim barev")
        self.color_combo.addItems(["Barevný", "Šedý", "ČB"])
        lbl_color.setBuddy(self.color_combo)
        scan_layout.addWidget(lbl_color)
        scan_layout.addWidget(self.color_combo)

        scan_group.setLayout(scan_layout)
        layout.addWidget(scan_group)

        # -- Tlačítka skenování --
        scan_btn_layout = QHBoxLayout()
        self.btn_scan = QPushButton("Skenovat stránku (Ctrl+S)")
        self.btn_scan.setAccessibleName("Skenovat stránku")
        self.btn_scan.clicked.connect(self.scan_pages)
        scan_btn_layout.addWidget(self.btn_scan)

        self.btn_scan_next = QPushButton("Skenovat další stránku (Ctrl+Shift+N)")
        self.btn_scan_next.setAccessibleName("Skenovat další stránku – přidat ke stávajícímu dokumentu")
        self.btn_scan_next.clicked.connect(self.scan_next_page)
        scan_btn_layout.addWidget(self.btn_scan_next)

        self.btn_scan_all = QPushButton("Skenovat vše (Ctrl+Shift+S)")
        self.btn_scan_all.setAccessibleName("Skenovat všechny stránky dávkově")
        self.btn_scan_all.clicked.connect(self.scan_all_pages)
        scan_btn_layout.addWidget(self.btn_scan_all)
        layout.addLayout(scan_btn_layout)

        # -- Seznam stránek --
        self.page_list = QListWidget()
        self.page_list.setAccessibleName("Seznam naskenovaných stránek")
        self.page_list.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        layout.addWidget(QLabel("Naskenované stránky:"))
        layout.addWidget(self.page_list)

        page_btn_layout = QHBoxLayout()
        btn_delete_page = QPushButton("Smazat stránku")
        btn_delete_page.setAccessibleName("Smazat vybranou stránku")
        btn_delete_page.clicked.connect(self.delete_page)
        page_btn_layout.addWidget(btn_delete_page)

        btn_clear_pages = QPushButton("Smazat vše")
        btn_clear_pages.setAccessibleName("Smazat všechny stránky")
        btn_clear_pages.clicked.connect(self.clear_pages)
        page_btn_layout.addWidget(btn_clear_pages)
        layout.addLayout(page_btn_layout)

        # -- Tlačítka OCR a uložení – logicky oddělené --
        ocr_btn_layout = QHBoxLayout()
        self.btn_ocr = QPushButton("Spustit OCR (Ctrl+O)")
        self.btn_ocr.setAccessibleName("Spustit OCR – rozpozná text ze stránek")
        self.btn_ocr.clicked.connect(self.run_ocr)
        self.btn_ocr.setEnabled(False)
        ocr_btn_layout.addWidget(self.btn_ocr)

        self.btn_save_document = QPushButton("Uložit dokument (Ctrl+U)")
        self.btn_save_document.setAccessibleName("Uložit dokument – uloží naskenované stránky nebo rozpoznaný text")
        self.btn_save_document.clicked.connect(self.save_document)
        self.btn_save_document.setEnabled(False)
        ocr_btn_layout.addWidget(self.btn_save_document)

        self.btn_read = QPushButton("Přečíst text (Ctrl+P)")
        self.btn_read.setAccessibleName("Přečíst naposledy rozpoznaný text")
        self.btn_read.clicked.connect(self.read_last_text)
        self.btn_read.setEnabled(False)
        ocr_btn_layout.addWidget(self.btn_read)
        layout.addLayout(ocr_btn_layout)

        btn_export_img = QPushButton("Exportovat obrázky (Ctrl+E)")
        btn_export_img.setAccessibleName("Exportovat naskenované obrázky")
        btn_export_img.clicked.connect(self.export_images)
        layout.addWidget(btn_export_img)

        # -- Uživatelská makra --
        macro_group = QGroupBox("Uživatelská makra")
        self._macro_container = QVBoxLayout()
        self._macro_buttons_layout = QVBoxLayout()
        self._macro_container.addLayout(self._macro_buttons_layout)
        btn_edit_macros = QPushButton("Spravovat makra...")
        btn_edit_macros.setAccessibleName("Otevřít editor maker")
        btn_edit_macros.clicked.connect(self._open_macro_editor)
        self._macro_container.addWidget(btn_edit_macros)
        macro_group.setLayout(self._macro_container)
        layout.addWidget(macro_group)

        scroll.setWidget(container)

        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.addWidget(scroll)
        self.setLayout(main_layout)
        self.setStyleSheet("""
            QPushButton, QComboBox {
                padding: 6px;
                min-height: 1.5em;
            }
            QCheckBox {
                padding: 6px;
                min-height: 1.75em;
                spacing: 8px;
            }
            QCheckBox::indicator {
                width: 18px;
                height: 18px;
            }
            QGroupBox {
                margin-top: 1.2em;
                padding-top: 8px;
            }
            QGroupBox::title {
                subcontrol-origin: margin;
                subcontrol-position: top left;
                padding: 0 4px;
            }
        """)
        self._rebuild_macro_buttons()

        # -- Tab order --
        self.setTabOrder(self.lang_combo, self.engine_combo)
        self.setTabOrder(self.engine_combo, self.preprocess_cb)
        self.setTabOrder(self.preprocess_cb, self.diacritics_cb)
        self.setTabOrder(self.diacritics_cb, self.batch_cb)
        self.setTabOrder(self.batch_cb, self.device_combo)
        self.setTabOrder(self.device_combo, self.source_combo)
        self.setTabOrder(self.source_combo, self.dpi_combo)
        self.setTabOrder(self.dpi_combo, self.color_combo)
        self.setTabOrder(self.color_combo, self.btn_scan)
        self.setTabOrder(self.btn_scan, self.btn_scan_next)
        self.setTabOrder(self.btn_scan_next, self.btn_scan_all)
        self.setTabOrder(self.btn_scan_all, self.page_list)
        self.setTabOrder(self.page_list, btn_delete_page)
        self.setTabOrder(btn_delete_page, btn_clear_pages)
        self.setTabOrder(btn_clear_pages, self.btn_ocr)
        self.setTabOrder(self.btn_ocr, self.btn_save_document)
        self.setTabOrder(self.btn_save_document, self.btn_read)
        self.setTabOrder(self.btn_read, btn_export_img)

        # Uložit reference pro testování Tab order a pro focus handling
        self._scroll_area = scroll
        self._scroll_container = container
        # QGroupBox nesmí být focusable – pojistka (default NoFocus, ale explicitně)
        for _gb in (ocr_group, scan_group, macro_group):
            _gb.setFocusPolicy(Qt.FocusPolicy.NoFocus)
        # Dekorativní QLabel také ne
        for _lbl in (lbl_lang, lbl_engine, lbl_diac, lbl_device, lbl_source, lbl_dpi, lbl_color):
            _lbl.setFocusPolicy(Qt.FocusPolicy.NoFocus)

    # ---------- Počáteční fokus pro NVDA ----------
    def showEvent(self, event) -> None:
        super().showEvent(event)
        if not self._initial_focus_done:
            # Odložit na konec event loopu – okno musí být viditelné a mapped.
            # QTimer.singleShot(0, ...) je robustnější než přímé setFocus v __init__.
            QTimer.singleShot(0, self._set_initial_focus)

    def _set_initial_focus(self) -> None:
        if self._initial_focus_done:
            return
        # Pokud okno ještě není aktivní (Windows activation delay), zkusit znovu
        # – neagresivně, bez krádeže fokusu. Max 5 pokusů po 100 ms.
        if not self.isActiveWindow():
            # Zjistit, zda už byl pokus; uložit counter do atributu
            tries = getattr(self, "_initial_focus_tries", 0)
            if tries < 5:
                self._initial_focus_tries = tries + 1
                QTimer.singleShot(100, self._set_initial_focus)
                return
            # Po vyčerpání pokusů nastavit fokus i bez isActiveWindow –
            # QWidget.setFocus() vyžaduje active window pro NVDA, ale
            # fokus bude doručen jakmile se okno aktivuje uživatelem/WM.
            # Nepoužívat QApplication.setActiveWindow() (deprecated) ani
            # agresivní activateWindow() pokud uživatel mezitím přešel jinam.
        # Nepoužívat zastaralé QApplication.setActiveWindow().
        # activateWindow() pouze pokud je to bezpečné a okno je viditelné,
        # ale nekrást fokus – proto jen pokud je okno viditelné a ještě neaktivní
        # a uživatel zjevně aplikaci právě spustil (první show).
        # V praxi stačí setFocus; Windows dá aktivaci automaticky při spuštění.
        self.lang_combo.setFocus(Qt.FocusReason.OtherFocusReason)
        # Qt interně pošle QAccessible::Focus event přes UIA bridge.
        # Explicitní QAccessible.updateAccessibility() by bylo duplicitní
        # a v PyQt6 6.11 není QAccessible v Python API vůbec exponován.
        # Pokus o podmíněné poslání pouze pokud by bylo dostupné:
        try:
            import importlib
            qacc = importlib.import_module("PyQt6.QtGui")
            QAccessible = getattr(qacc, "QAccessible", None)
            if QAccessible is not None and hasattr(QAccessible, "updateAccessibility"):
                # isActive() guard – posílat jen když běží AT
                is_active = getattr(QAccessible, "isActive", None)
                if is_active is None or is_active():
                    QAccessible.updateAccessibility(
                        self.lang_combo, 0, QAccessible.Event.Focus
                    )
        except Exception:
            pass
        self._initial_focus_done = True

    def _setup_shortcuts(self) -> None:
        QShortcut(QKeySequence("Ctrl+S"), self).activated.connect(self.scan_pages)
        QShortcut(QKeySequence("Ctrl+Shift+N"), self).activated.connect(self.scan_next_page)
        QShortcut(QKeySequence("Ctrl+Shift+S"), self).activated.connect(self.scan_all_pages)
        QShortcut(QKeySequence("Ctrl+O"), self).activated.connect(self.run_ocr)
        QShortcut(QKeySequence("Ctrl+U"), self).activated.connect(self.save_document)
        QShortcut(QKeySequence("Ctrl+P"), self).activated.connect(self.read_last_text)
        QShortcut(QKeySequence("Ctrl+E"), self).activated.connect(self.export_images)
        QShortcut(QKeySequence("Ctrl+Q"), self).activated.connect(self.close)
        QShortcut(QKeySequence("Delete"), self).activated.connect(self.delete_page)
        QShortcut(QKeySequence("Ctrl+M"), self).activated.connect(self._open_macro_editor)
        self._setup_macro_shortcuts()

    def _setup_macro_shortcuts(self) -> None:
        for sc in self._macro_shortcuts:
            sc.setEnabled(False)
            sc.deleteLater()
        self._macro_shortcuts = []
        for macro in self.macro_manager.macros:
            if macro.shortcut:
                sc = QShortcut(QKeySequence(macro.shortcut), self)
                sc.activated.connect(lambda checked=False, m=macro: self._run_macro(m))
                self._macro_shortcuts.append(sc)

    # ---------- Settings persistence ----------
    def _load_settings(self) -> None:
        lang_idx = self.settings.value("ocr/lang_index", 0, type=int)
        if 0 <= lang_idx < self.lang_combo.count():
            self.lang_combo.setCurrentIndex(lang_idx)
        engine_idx = self.settings.value("ocr/engine_index", 0, type=int)
        if 0 <= engine_idx < self.engine_combo.count():
            self.engine_combo.setCurrentIndex(engine_idx)
        dpi_val = self.settings.value("scan/dpi", 300, type=int)
        try:
            dpi_val = int(dpi_val)
        except (TypeError, ValueError):
            dpi_val = 300
        _allowed_dpi = {200, 300, 400, 600, 1200}
        if dpi_val not in _allowed_dpi:
            dpi_val = 300
        idx = self.dpi_combo.findData(dpi_val)
        if idx >= 0:
            self.dpi_combo.setCurrentIndex(idx)
        else:
            idx_def = self.dpi_combo.findData(300)
            if idx_def >= 0:
                self.dpi_combo.setCurrentIndex(idx_def)
        source_idx = self.settings.value("scan/source_index", 0, type=int)
        if 0 <= source_idx < self.source_combo.count():
            self.source_combo.setCurrentIndex(source_idx)
        color_idx = self.settings.value("scan/color_index", 0, type=int)
        if 0 <= color_idx < self.color_combo.count():
            self.color_combo.setCurrentIndex(color_idx)
        self.batch_cb.setChecked(self.settings.value("scan/batch", False, type=bool))
        self.preprocess_cb.setChecked(self.settings.value("ocr/preprocess", True, type=bool))
        self.diacritics_cb.setChecked(self.settings.value("ocr/diacritics", False, type=bool))

    def _on_diacritics_toggled(self, checked: bool) -> None:
        # Hlasová odezva při přepnutí
        if checked:
            speak("Oprava české diakritiky zapnuta.")
        else:
            speak("Oprava české diakritiky vypnuta.")
        self._save_settings()

    def _save_settings(self) -> None:
        self.settings.setValue("ocr/lang_index", self.lang_combo.currentIndex())
        self.settings.setValue("ocr/engine_index", self.engine_combo.currentIndex())
        self.settings.setValue("scan/dpi", self.dpi_combo.currentData())
        self.settings.setValue("scan/source_index", self.source_combo.currentIndex())
        self.settings.setValue("scan/color_index", self.color_combo.currentIndex())
        self.settings.setValue("scan/batch", self.batch_cb.isChecked())
        self.settings.setValue("ocr/preprocess", self.preprocess_cb.isChecked())
        self.settings.setValue("ocr/diacritics", self.diacritics_cb.isChecked())

    def closeEvent(self, event) -> None:
        self._save_settings()
        super().closeEvent(event)

    # ---------- Device management ----------
    def reload_devices(self) -> None:
        try:
            devices = self.scanner.list_devices()
            self.device_combo.clear()
            for d in devices:
                self.device_combo.addItem(d)
            if devices:
                self.scanner.connect_device(devices[0])
            elif not self.scanner.naps2_exe:
                QMessageBox.warning(
                    self, "NAPS2 nenalezen",
                    "NAPS2.Console.exe nebyl nalezen. Ujistěte se, že je NAPS2 nainstalován."
                )
        except Exception as e:
            QMessageBox.critical(self, "Chyba", f"Nenašel jsem žádný skener:\n{e}")

    # ---------- Scanning ----------
    def scan_pages(self) -> None:
        # If pages already exist, never clear without confirmation.
        if self.scanned_images:
            count = len(self.scanned_images)
            speak(f"Máte již {count} naskenovaných stránek. Potvrďte smazání nebo přidejte další stránku.")
            msg = QMessageBox(self)
            msg.setWindowTitle("Nový dokument")
            msg.setText(f"Máte již {count} naskenovaných stránek. Chcete začít nový dokument?")
            msg.setInformativeText("Stávající stránky budou smazány. Můžete také přidat další stránku ke stávajícímu dokumentu.")
            btn_new = msg.addButton("Ano, smazat a začít znovu", QMessageBox.ButtonRole.YesRole)
            btn_new.setAccessibleName("Ano, smazat stávající stránky a začít nový dokument")
            btn_append = msg.addButton("Přidat další stránku", QMessageBox.ButtonRole.ActionRole)
            btn_append.setAccessibleName("Přidat další stránku ke stávajícímu dokumentu")
            btn_cancel = msg.addButton("Zrušit", QMessageBox.ButtonRole.NoRole)
            btn_cancel.setAccessibleName("Zrušit, zachovat stránky")
            msg.setDefaultButton(btn_cancel)
            msg.exec()
            clicked = msg.clickedButton()
            if clicked == btn_append:
                self.scan_next_page()
                return
            if clicked != btn_new:
                speak("Skenování zrušeno, stránky zachovány.")
                self.page_list.setFocus()
                return
            # User confirmed new document – clear everything
            self.scanned_images.clear()
            self.raw_scanned_images.clear()
            self.page_list.clear()
            self.last_ocr_results.clear()
            self._last_text = ""
            self._last_original_text = ""
            self._last_edited_pages = None
            self._diacritics_failed = False
            self.btn_ocr.setEnabled(False)
            self.btn_read.setEnabled(False)
            self._update_save_button_state()
        else:
            # No pages – ensure clean state (in case of residual OCR results)
            self.last_ocr_results.clear()
            self._last_text = ""
            self._last_original_text = ""
            self._last_edited_pages = None
            self._diacritics_failed = False

        speak("Zahajuji skenování.")
        # Keep button states disabled until at least one page succeeds
        self.btn_ocr.setEnabled(False)
        self.btn_read.setEnabled(False)
        self._update_save_button_state()

        # Remember count before scanning for correct voice after
        start_count = len(self.scanned_images)

        while True:
            self.last_scanned_image = None
            self._scan_cancelled = False
            self._do_scan_single()

            if self._scan_cancelled or self.last_scanned_image is None:
                break

            if self.batch_cb.isChecked():
                continue

            speak("Další stránka?")
            msg = QMessageBox(self)
            msg.setWindowTitle("Další stránka")
            msg.setText("Další stránka?")
            btn_yes = msg.addButton("Ano", QMessageBox.ButtonRole.YesRole)
            btn_yes.setAccessibleName("Ano, skenovat další stránku")
            btn_no = msg.addButton("Ne", QMessageBox.ButtonRole.NoRole)
            btn_no.setAccessibleName("Ne, ukončit skenování")
            msg.exec()
            if msg.clickedButton() != btn_yes:
                break

        if self.scanned_images:
            self.btn_ocr.setEnabled(True)
            self._update_save_button_state()
            # In new-document mode we cleared page_list, so repopulate.
            # If start_count was 0 we are in new mode; otherwise we already cleared.
            self.page_list.clear()
            self.page_list.addItems([f"Stránka {i+1}" for i in range(len(self.scanned_images))])
            self.page_list.setFocus()
            speak(f"Skenování dokončeno. {len(self.scanned_images)} stránek.")
        else:
            self.btn_scan.setFocus()
            self._update_save_button_state()
            speak("Skenování dokončeno, žádné stránky.")

    def scan_next_page(self) -> None:
        speak("Přidávám další stránku ke stávajícímu dokumentu.")
        count_before = len(self.scanned_images)
        self.last_scanned_image = None
        self._scan_cancelled = False
        self._do_scan_single()

        if self._scan_cancelled or self.last_scanned_image is None:
            speak("Skenování další stránky zrušeno.")
            if self.scanned_images:
                self.page_list.setFocus()
            else:
                self.btn_scan.setFocus()
            return

        # _scan_finished already appended images; now update page_list incrementally
        new_num = len(self.scanned_images)
        # Guard against double-add if scan_pages bulk path was used – but scan_next_page
        # never clears, so page_list should have count_before items
        if self.page_list.count() < new_num:
            self.page_list.addItem(f"Stránka {new_num}")
        else:
            # Fallback: rebuild to stay consistent
            self.page_list.clear()
            self.page_list.addItems([f"Stránka {i+1}" for i in range(new_num)])

        self.btn_ocr.setEnabled(True)
        self._update_save_button_state()
        self.page_list.setCurrentRow(new_num - 1)
        self.page_list.setFocus()
        total = len(self.scanned_images)
        # Keep existing OCR results for pages 1..count_before; new page is not OCRed yet
        if self.last_ocr_results and len(self.last_ocr_results) < total:
            speak(f"Stránka {new_num} přidána. Celkem {total} stránek. OCR nové stránky zatím nebylo provedeno.")
        else:
            speak(f"Stránka {new_num} přidána. Celkem {total} stránek.")

    def _renumber_pages(self) -> None:
        for i in range(self.page_list.count()):
            self.page_list.item(i).setText(f"Stránka {i+1}")

    # ---------- Helpers pro stav OCR a tlačítko Uložit ----------
    def _has_ocr_text(self) -> bool:
        for page in self.last_ocr_results:
            for item in page:
                if item.text and item.text.strip():
                    return True
        # fallback – _last_text může obsahovat text i bez bbox
        if self._last_text and self._last_text.strip():
            # ověř že text není jen headery
            stripped = re.sub(r"---\s*Stránka\s+\d+\s*---", "", self._last_text).strip()
            return bool(stripped)
        return False

    def _is_ocr_empty(self) -> bool:
        return not self.last_ocr_results or not self._has_ocr_text()

    def _is_ocr_complete(self) -> bool:
        total = len(self.scanned_images)
        if total == 0:
            return False
        if len(self.last_ocr_results) != total:
            return False
        return self._has_ocr_text()

    def _missing_ocr_count(self) -> int:
        return max(0, len(self.scanned_images) - len(self.last_ocr_results))

    def _update_save_button_state(self) -> None:
        has_pages = len(self.scanned_images) > 0
        if hasattr(self, "btn_save_document"):
            self.btn_save_document.setEnabled(has_pages)

    def scan_all_pages(self) -> None:
        old_batch = self.batch_cb.isChecked()
        self.batch_cb.setChecked(True)
        self.scan_pages()
        self.batch_cb.setChecked(old_batch)

    def _do_scan_single(self) -> None:
        device_name = self.device_combo.currentText()
        if not device_name:
            QMessageBox.warning(self, "Chyba", "Není vybrán žádný skener.")
            return

        self.scanner.connect_device(device_name)
        dpi = self.dpi_combo.currentData()
        if dpi is None:
            dpi = 300
        source = self.source_combo.currentText()
        color_text = self.color_combo.currentText()
        color_mode = (
            "Color" if color_text == "Barevný"
            else "Grayscale" if color_text == "Šedý"
            else "BlackWhite"
        )

        self.progress_dialog = QProgressDialog(
            "Skenuji (přes NAPS2)...", "Zrušit", 0, 0, self
        )
        self.progress_dialog.canceled.connect(self._cancel_scan)
        self.progress_dialog.show()

        self.scan_thread = ScanThread(self.scanner, dpi, color_mode, source)
        self.scan_thread.finished.connect(self._scan_finished)
        self.scan_thread.error.connect(self._scan_error)
        self.scan_thread.start()

        loop = QEventLoop()
        self.scan_thread.finished.connect(loop.quit)
        self.scan_thread.error.connect(loop.quit)
        loop.exec()

        self.progress_dialog.cancel()

    def _cancel_scan(self) -> None:
        self._scan_cancelled = True
        speak("Skenování zrušeno.")

    def _scan_finished(self, img: Image.Image) -> None:
        self.progress_dialog.cancel()
        if self._scan_cancelled:
            self.last_scanned_image = None
            return

        preview = PreviewDialog(img, len(self.scanned_images) + 1)
        if preview.exec():
            original = preview.image
            processed = original
            if self.preprocess_cb.isChecked():
                processed = preprocess_image(original)
            self.raw_scanned_images.append(original)
            self.scanned_images.append(processed)
            self.last_scanned_image = processed
        else:
            self.last_scanned_image = None

    def _scan_error(self, msg: str) -> None:
        self.progress_dialog.cancel()
        speak(f"Chyba skenování: {msg}")
        QMessageBox.critical(self, "Chyba skenování", msg)
        self.last_scanned_image = None

    # ---------- Page management ----------
    def delete_page(self) -> None:
        row = self.page_list.currentRow()
        if row < 0:
            speak("Není vybrána žádná stránka k smazání.")
            return
        page_num = row + 1
        speak(f"Potvrdit smazání stránky {page_num}.")
        msg = QMessageBox(self)
        msg.setWindowTitle("Potvrdit smazání stránky")
        msg.setText(f"Opravdu chcete smazat stránku {page_num}?")
        msg.setInformativeText("Tuto akci nelze vrátit.")
        btn_yes = msg.addButton("Ano", QMessageBox.ButtonRole.YesRole)
        btn_yes.setAccessibleName(f"Ano, smazat stránku {page_num}")
        btn_no = msg.addButton("Ne", QMessageBox.ButtonRole.NoRole)
        btn_no.setAccessibleName("Ne, ponechat stránku")
        msg.setDefaultButton(btn_no)
        msg.exec()
        if msg.clickedButton() != btn_yes:
            speak("Smazání zrušeno.")
            self.page_list.setFocus()
            self.page_list.setCurrentRow(row)
            return

        self.page_list.takeItem(row)
        del self.scanned_images[row]
        del self.raw_scanned_images[row]
        if self.last_ocr_results and row < len(self.last_ocr_results):
            del self.last_ocr_results[row]
        if self._last_edited_pages is not None and row < len(self._last_edited_pages):
            del self._last_edited_pages[row]
        # Pokud byl _last_text s headery, přegeneruj ho podle zbývajících stránek
        if self._last_edited_pages is not None:
            # Přegeneruj _last_text aby reflektoval smazání stránky
            parts = []
            for idx, p in enumerate(self._last_edited_pages):
                parts.append(f"--- Stránka {idx+1} ---\n{p}\n\n")
            self._last_text = "".join(parts)
        # Renumber remaining items to keep Stránka 1..N consistent
        self._renumber_pages()
        remaining = len(self.scanned_images)
        if remaining == 0:
            self.btn_ocr.setEnabled(False)
            self.btn_read.setEnabled(False)
            self._last_text = ""
            self._last_original_text = ""
            self._last_edited_pages = None
            self._diacritics_failed = False
            self._update_save_button_state()
            speak("Všechny stránky smazány.")
            self.btn_scan.setFocus()
        else:
            speak(f"Stránka smazána. Zbývá {remaining} stránek.")
            self._update_save_button_state()
            # Focus na stejnou pozici nebo poslední
            next_row = min(row, remaining - 1)
            self.page_list.setCurrentRow(next_row)
            self.page_list.setFocus()

    def clear_pages(self) -> None:
        count = len(self.scanned_images)
        if count == 0:
            speak("Žádné stránky ke smazání.")
            return
        speak(f"Potvrdit smazání všech {count} stránek.")
        msg = QMessageBox(self)
        msg.setWindowTitle("Potvrdit smazání všech stránek")
        msg.setText(f"Opravdu chcete smazat všech {count} naskenovaných stránek?")
        msg.setInformativeText("Budou odstraněny všechny naskenované stránky a jejich OCR výsledky. Tuto akci nelze vrátit.")
        btn_yes = msg.addButton("Ano", QMessageBox.ButtonRole.YesRole)
        btn_yes.setAccessibleName("Ano, smazat všechny stránky")
        btn_no = msg.addButton("Ne", QMessageBox.ButtonRole.NoRole)
        btn_no.setAccessibleName("Ne, ponechat stránky")
        msg.setDefaultButton(btn_no)
        msg.exec()
        if msg.clickedButton() != btn_yes:
            speak("Smazání zrušeno.")
            self.page_list.setFocus()
            return
        self.page_list.clear()
        self.scanned_images.clear()
        self.raw_scanned_images.clear()
        self.last_ocr_results.clear()
        self._last_text = ""
        self._last_original_text = ""
        self._last_edited_pages = None
        self._diacritics_failed = False
        self.btn_ocr.setEnabled(False)
        self.btn_read.setEnabled(False)
        self._update_save_button_state()
        speak(f"Všech {count} stránek smazáno.")
        self.btn_scan.setFocus()

    # ---------- OCR ----------
    def _run_ocr_incremental(self, start_idx: int) -> None:
        """OCR pouze pro nově přidané stránky start_idx..end, zachová 1..start_idx-1."""
        new_images = self.scanned_images[start_idx:]
        if not new_images:
            return
        total_new = len(new_images)
        total_all = len(self.scanned_images)
        self._ocr_total_pages = total_all
        self._ocr_last_announced_page = start_idx  # aby progress hlásil od start_idx+1
        speak(f"Zahajuji OCR {total_new} nových stránek (celkem {total_all}).")
        base_engine = self.engine_combo.currentText()
        base_lang = self.lang_combo.currentText()
        # Jednorázový pokus pro nové stránky – stačí base pokus, fallback ponechán pro full flow
        # Zde použijeme stejný mechanismus jako _run_ocr_once ale s merge
        old_results = list(self.last_ocr_results)
        old_text = self._last_text

        # Dočasné sloty pro merge
        new_results_holder: list[list[OcrResult]] = []
        new_text_holder: list[str] = []

        loop = QEventLoop()

        def on_results(r):
            new_results_holder.clear()
            new_results_holder.extend(r)

        def on_finished(txt: str):
            new_text_holder.append(txt)
            loop.quit()

        # Vytvoř thread pro nové stránky
        diac_enabled = self.diacritics_cb.isChecked() if hasattr(self, "diacritics_cb") else False
        self._diacritics_failed = False
        self._diacritics_enabled_cache = diac_enabled
        self.ocr_thread = create_ocr_thread(base_engine, new_images, base_lang, diacritics_enabled=diac_enabled)
        self.ocr_thread.ocr_results.connect(on_results)
        self.ocr_thread.finished.connect(on_finished)
        if hasattr(self.ocr_thread, "diacritics_failed"):
            self.ocr_thread.diacritics_failed.connect(self._on_diacritics_failed)
        if base_engine == "EasyOCR":
            self.progress_dialog = QProgressDialog("Načítám EasyOCR model...", "Zrušit", 0, 0, self)
            self.ocr_thread.model_loading.connect(self._on_model_loaded)
        else:
            self.progress_dialog = QProgressDialog("Probíhá OCR (Tesseract)...", "Zrušit", 0, 100, self)
        self.ocr_thread.progress.connect(self._on_ocr_progress)
        self.progress_dialog.show()
        self.ocr_thread.start()
        loop.exec()
        self.progress_dialog.cancel()

        if not new_results_holder:
            speak("OCR nových stránek neprodukovalo žádné výsledky.")
            return
        # Oprav číslování stránek v novém textu (thread čísluje od 1)
        # Sestav korektní full_text pro nové stránky s posunutým číslováním
        corrected_new_text = ""
        for i, page_results in enumerate(new_results_holder):
            page_num = start_idx + i + 1
            # Extrahuj text stránky z new_text_holder[0] – jednodušší rekonstruovat z page_results
            # ale zachovej původní diakritizovaný text včetně odřádkování – použij display_text
            page_text_parts = [r.display_text for r in page_results]
            # Pokud page_results obsahuje věty s newline, spoj mezery
            page_text = "\n".join(page_text_parts)
            # Pokud původní new_text_holder obsahuje více, použij ho přímo s přečíslováním
            # Pro jednoduchost použij page_text z results
            corrected_new_text += f"--- Stránka {page_num} ---\n{page_text}\n\n"

        # Merge
        self.last_ocr_results = old_results + new_results_holder
        # Pokud byl původní full_text s headery, připoj nové; jinak použij corrected
        # Pro zachování přesnosti použij starý text + corrected_new_text
        if old_text and old_text.strip():
            self._last_text = old_text.rstrip() + "\n\n" + corrected_new_text
        else:
            self._last_text = corrected_new_text
        # Aktualizuj per-page
        try:
            pages = _split_text_by_page_headers(self._last_text)
            if pages is not None and len(pages) == len(self.last_ocr_results):
                self._last_edited_pages = pages
            else:
                self._last_edited_pages = None
        except Exception:
            self._last_edited_pages = None
        # Rekonstrukce originálu
        try:
            orig_parts = []
            for idx, page in enumerate(self.last_ocr_results):
                if page:
                    texts = [r.original_text if r.original_text is not None else r.text for r in page]
                    orig_parts.append(f"--- Stránka {idx + 1} ---\n" + "\n".join(texts) + "\n\n")
            self._last_original_text = "".join(orig_parts) if any(orig_parts) else self._last_text
        except Exception:
            self._last_original_text = self._last_text
        self.btn_read.setEnabled(True)
        # Hlasová odezva
        from diacritics import should_diacritize
        lang = self.lang_combo.currentText() if hasattr(self, "lang_combo") else ""
        diac_active = should_diacritize(lang, self._diacritics_enabled_cache)
        if self._diacritics_failed and diac_active:
            speak("OCR dokončeno. Oprava české diakritiky se nepodařila. Byl zachován původní text.")
        elif diac_active:
            speak("OCR dokončeno a česká diakritika opravena.")
        else:
            speak(f"OCR dokončeno pro {total_all} stránek.")

        self._ocr_mode = "interactive"
        # Pokud je pending save (voláno z Uložit dokument), náhled nezobrazuj – rovnou pokračuj k uložení
        if getattr(self, "_pending_save_format", None):
            return
        self._show_ocr_preview(self._last_text)

    def run_ocr(self) -> None:
        if not self.scanned_images:
            QMessageBox.warning(self, "Chyba", "Nejdříve naskenujte stránky.")
            return
        total = len(self.scanned_images)
        # Inkrementální OCR: pokud již máme OCR pro část stránek, zpracuj pouze nové
        if self.last_ocr_results and 0 < len(self.last_ocr_results) < total:
            # Zachovej pořadí 1..N, přidej pouze chybějící N+1..total
            self._run_ocr_incremental(len(self.last_ocr_results))
            return
        self._ocr_total_pages = total
        self._ocr_last_announced_page = 0
        speak(f"Zahajuji OCR {total} stránek.")
        self._run_ocr_flow(interactive=True)

    def _start_ocr(
        self,
        images: list[Image.Image],
        lang: str,
        engine: str,
        on_done: Optional[Callable[[], None]] = None,
    ) -> None:
        diac_enabled = self.diacritics_cb.isChecked() if hasattr(self, "diacritics_cb") else False
        self._diacritics_failed = False
        self._diacritics_enabled_cache = diac_enabled
        self.ocr_thread = create_ocr_thread(engine, images, lang, diacritics_enabled=diac_enabled)
        self.ocr_thread.finished.connect(self._ocr_finished)
        self.ocr_thread.ocr_results.connect(self._set_ocr_results)
        if hasattr(self.ocr_thread, "diacritics_failed"):
            self.ocr_thread.diacritics_failed.connect(self._on_diacritics_failed)
        if on_done:
            self.ocr_thread.finished.connect(on_done)

        if engine == "EasyOCR":
            self.progress_dialog = QProgressDialog(
                "Načítám EasyOCR model...", "Zrušit", 0, 0, self
            )
            self.ocr_thread.model_loading.connect(self._on_model_loaded)
        else:
            self.progress_dialog = QProgressDialog(
                "Probíhá OCR (Tesseract)...", "Zrušit", 0, 100, self
            )

        self.ocr_thread.progress.connect(self._on_ocr_progress)
        self.progress_dialog.show()
        self.ocr_thread.start()

    def _build_attempts(self) -> list[tuple[str, str, list[Image.Image]]]:
        base_engine = self.engine_combo.currentText()
        base_lang = self.lang_combo.currentText()
        alt_engine = "EasyOCR" if base_engine == "Tesseract" else "Tesseract"

        attempts: list[tuple[str, str, list[Image.Image]]] = [
            (base_engine, base_lang, list(self.scanned_images)),
            (alt_engine, base_lang, list(self.scanned_images)),
        ]

        for alt_lang in [
            "eng (Angličtina)", "ces (Čeština)", "deu (Němčina)",
            "fra (Francouzština)", "ita (Italština)", "pol (Polština)",
        ]:
            if alt_lang != base_lang:
                attempts.append((base_engine, alt_lang, list(self.scanned_images)))
                attempts.append((alt_engine, alt_lang, list(self.scanned_images)))

        alt_images = [alt_preprocess_image(r) for r in self.raw_scanned_images]
        if alt_images:
            attempts.append((base_engine, base_lang, alt_images))
            attempts.append((alt_engine, base_lang, alt_images))
        return attempts

    def _run_ocr_once(
        self, lang: str, engine: str, images: list[Image.Image]
    ) -> None:
        loop = QEventLoop()
        self._start_ocr(images, lang, engine, on_done=loop.quit)
        loop.exec()

    def _run_ocr_flow(self, interactive: bool) -> None:
        attempts = self._build_attempts()
        total = len(attempts)
        found = False
        for i, (engine, lang, images) in enumerate(attempts, start=1):
            if self._scan_cancelled:
                break
            speak(f"OCR pokus {i} z {total}.")
            self._ocr_mode = "macro"
            self._run_ocr_once(lang, engine, images)
            if self._has_ocr_text():
                found = True
                break

        if interactive:
            if found:
                self._ocr_mode = "interactive"
                # pending save nesmí zobrazit druhý náhled
                if not getattr(self, "_pending_save_format", None):
                    self._show_ocr_preview(self._last_text)
            else:
                action = self._ask_no_text_action()
                if action == "rescan":
                    speak("Znovu skenuji stránky a opakuji OCR.")
                    self.scan_pages()
                    if self.scanned_images:
                        self._run_ocr_flow(interactive=True)
                elif action == "continue":
                    self._ocr_mode = "interactive"
                    if not getattr(self, "_pending_save_format", None):
                        self._show_ocr_preview(self._last_text)

    def _ask_no_text_action(self) -> str:
        speak("Nebyl rozpoznán žádný text.")
        msg = QMessageBox(self)
        msg.setWindowTitle("Nerozpoznán žádný text")
        msg.setText("Po všech pokusech nebyl rozpoznán žádný text.")
        msg.setInformativeText("Co chcete udělat?")
        btn_rescan = msg.addButton("Znovu naskenovat", QMessageBox.ButtonRole.YesRole)
        btn_rescan.setAccessibleName("Znovu naskenovat stránky")
        btn_cont = msg.addButton("Pokračovat k uložení", QMessageBox.ButtonRole.ActionRole)
        btn_cont.setAccessibleName("Pokračovat k náhledu a uložení")
        btn_cancel = msg.addButton("Zrušit", QMessageBox.ButtonRole.NoRole)
        btn_cancel.setAccessibleName("Zrušit")
        msg.exec()
        clicked = msg.clickedButton()
        if clicked == btn_rescan:
            return "rescan"
        if clicked == btn_cont:
            return "continue"
        return "cancel"

    def _execute_ocr(self) -> None:
        # Ensure per-page announcements also work for macro/non-interactive path
        self._ocr_total_pages = len(self.scanned_images)
        self._ocr_last_announced_page = 0
        self._run_ocr_flow(interactive=False)

    def _on_ocr_progress(self, value: int) -> None:
        self.progress_dialog.setValue(value)
        # Per-page voice feedback – throttled to once per page
        total = getattr(self, "_ocr_total_pages", len(self.scanned_images)) or 1
        # Derive current page from progress percent
        current_page = int(round(value / 100 * total)) if total else 0
        if current_page > 0 and current_page != getattr(self, "_ocr_last_announced_page", 0):
            # Avoid double-announce at 100 which is handled separately
            if value != 100:
                speak(f"OCR stránky {current_page} z {total}.")
            self._ocr_last_announced_page = current_page
        if value == 100:
            speak("OCR dokončeno")

    def _on_model_loaded(self, value: int) -> None:
        if value == 100:
            self.progress_dialog.setMaximum(100)
            self.progress_dialog.setLabelText("Probíhá OCR (EasyOCR)...")
            speak("EasyOCR model načten, zahajuji rozpoznávání")

    def _set_ocr_results(self, results: list[list[OcrResult]]) -> None:
        self.last_ocr_results = results

    def _on_diacritics_failed(self, msg: str) -> None:
        self._diacritics_failed = True

    def _ocr_finished(self, full_text: str) -> None:
        self.progress_dialog.cancel()
        self._last_text = full_text
        # Inicializuj per-page z full_text
        try:
            pages = _split_text_by_page_headers(full_text)
            if pages is not None and len(pages) == len(self.scanned_images):
                self._last_edited_pages = pages
            else:
                self._last_edited_pages = None
        except Exception:
            self._last_edited_pages = None
        # Rekonstrukce originálu z OcrResult.original_text pro zachování původního výsledku
        try:
            orig_parts = []
            for idx, page in enumerate(self.last_ocr_results):
                # pokud máme bbox výsledky, poskládej originál, jinak použij full_text
                if page:
                    texts = [r.original_text if r.original_text is not None else r.text for r in page]
                    orig_parts.append(f"--- Stránka {idx + 1} ---\n" + "\n".join(texts) + "\n\n")
                else:
                    orig_parts.append("")
            self._last_original_text = "".join(orig_parts) if any(orig_parts) else full_text
        except Exception:
            self._last_original_text = full_text
        self.btn_read.setEnabled(True)
        total = getattr(self, "_ocr_total_pages", len(self.scanned_images))
        # Hlasová odezva dle diakritiky
        from diacritics import should_diacritize
        lang = self.lang_combo.currentText() if hasattr(self, "lang_combo") else ""
        diac_active = should_diacritize(lang, self._diacritics_enabled_cache)
        if self._diacritics_failed and diac_active:
            speak("OCR dokončeno. Oprava české diakritiky se nepodařila. Byl zachován původní text.")
        elif diac_active:
            speak("OCR dokončeno a česká diakritika opravena.")
            if total:
                speak(f"OCR dokončeno pro {total} stránek.")
        else:
            if total:
                speak(f"OCR dokončeno pro {total} stránek.")
        # Pokud běží pending save (OCR spuštěno z dialogu Uložit), nepřekrývej náhledem
        if getattr(self, "_pending_save_format", None):
            return
        if self._ocr_mode == "interactive":
            self._show_ocr_preview(full_text)

    def _show_ocr_preview(self, full_text: str) -> None:
        """Zobrazí náhled OCR textu k editaci – SAMOSTATNÝ krok OCR, bez ukládání."""
        speak("Zobrazuji náhled rozpoznaného textu. Můžete jej upravit nebo přečíst. Uložení provedete tlačítkem Uložit dokument.")
        dialog = OcrPreviewDialog(full_text, self)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            self.btn_read.setFocus()
            return
        edited_text = dialog.get_text()
        self._last_text = edited_text
        try:
            pages = _split_text_by_page_headers(edited_text)
            if pages is not None:
                self._last_edited_pages = pages
            else:
                if len(self.scanned_images) == 1:
                    self._last_edited_pages = [edited_text]
                else:
                    self._last_edited_pages = None
        except Exception:
            self._last_edited_pages = None
        # Po náhledu nabídni uložení přes samostatné tlačítko, ale ne automaticky
        # (uživatel stiskne Uložit dokument)
        self.btn_save_document.setFocus()
        speak("Náhled uložen. Pro uložení stiskněte Uložit dokument.")

    def _ocr_show_save_ui(self, full_text: str) -> None:
        """Legacy wrapper – zachován pro makra. Nově jen zobrazí náhled bez auto-uložení."""
        self._show_ocr_preview(full_text)
        # Původní auto-uložení odstraněno – oddělení OCR a ukládání.

    # ---------- Uložit dokument – dvoustupňový tok ----------
    def save_document(self) -> None:
        """Hlavní vstup pro tlačítko Uložit dokument – nikdy nespouští OCR automaticky."""
        if not self.scanned_images:
            QMessageBox.warning(self, "Chyba", "Nejdříve naskenujte stránky.")
            speak("Žádné stránky k uložení.")
            return
        speak("Otevírám dialog Co chcete uložit.")
        kind_dlg = ChooseSaveKindDialog(self)
        if kind_dlg.exec() != QDialog.DialogCode.Accepted:
            self.page_list.setFocus()
            return
        kind = kind_dlg.selected_kind()
        if kind == "images":
            self._save_images_only_flow()
        else:  # ocr
            fmt_dlg = ChooseOcrFormatDialog(self)
            if fmt_dlg.exec() != QDialog.DialogCode.Accepted:
                self.page_list.setFocus()
                return
            fmt = fmt_dlg.selected_format()  # txt / docx / pdf_ocr
            self._pending_save_format = fmt
            try:
                self._try_save_with_ocr(fmt)
            finally:
                # pokud nebyl spuštěn OCR, vyčisti hned; pokud byl pending save úspěšný, už je vyčištěn
                if self._pending_save_format == fmt:
                    # nebyl spuštěn OCR, nebo selhal – vyčisti
                    self._pending_save_format = None

    def _save_images_only_flow(self) -> None:
        """PDF bez OCR – používá přímo self.scanned_images, nikdy OCR."""
        speak("Ukládám PDF bez OCR – pouze obrázky.")
        default_dir = QStandardPaths.writableLocation(QStandardPaths.StandardLocation.DocumentsLocation)
        if not default_dir:
            default_dir = os.path.expanduser("~")
        default_path = os.path.join(default_dir, "naskenovany_dokument.pdf")
        path, _ = QFileDialog.getSaveFileName(self, "Uložit dokument", default_path, "PDF – naskenované stránky (*.pdf)")
        if not path:
            self.page_list.setFocus()
            return
        if not path.lower().endswith(".pdf"):
            base, ext = os.path.splitext(path)
            if ext.lower() not in (".txt", ".pdf", ".docx"):
                path = path + ".pdf"
            elif ext.lower() != ".pdf":
                path = base + ".pdf"
        if os.path.exists(path):
            if not self._confirm_overwrite(path):
                speak("Ukládání zrušeno.")
                self.page_list.setFocus()
                return
        try:
            self._save_pdf_without_ocr(path)
        except Exception as e:
            QMessageBox.critical(self, "Chyba při ukládání", f"Nepodařilo se uložit soubor:\n{e}")
            speak("Chyba při ukládání souboru.")
            self.page_list.setFocus()

    def _try_save_with_ocr(self, fmt: str) -> None:
        total = len(self.scanned_images)
        missing = self._missing_ocr_count()
        is_empty = self._is_ocr_empty()
        is_complete = self._is_ocr_complete()
        # fmt: txt, docx, pdf_ocr
        if is_complete:
            self._do_save_ocr_format(fmt)
            self._pending_save_format = None
            return
        # parciální nebo prázdné – zobraz dialog
        if is_empty:
            action = self._ask_ocr_not_available(fmt)
        else:
            action = self._ask_partial_ocr(missing, total, fmt)
        if action == "run_ocr":
            # Spustit OCR a po dokončení automaticky pokračovat
            self._run_ocr_for_pending_save(fmt)
        elif action == "save_without" and fmt == "pdf_ocr":
            # Uživatel chce uložit bez OCR místo s OCR
            self._pending_save_format = None
            self._save_images_only_flow()
        elif action == "cancel":
            self._pending_save_format = None
            speak("Ukládání zrušeno.")
            self.page_list.setFocus()
        else:
            # pro txt/docx není save_without – jen cancel/run
            self._pending_save_format = None
            speak("Ukládání zrušeno.")
            self.page_list.setFocus()

    def _ask_ocr_not_available(self, fmt: str) -> str:
        """Dialog OCR není k dispozici. Vrátí run_ocr / save_without / cancel."""
        speak("OCR není k dispozici.")
        fmt_human = {"txt": "TXT", "docx": "DOCX", "pdf_ocr": "PDF s OCR"}[fmt]
        msg = QMessageBox(self)
        msg.setWindowTitle("OCR není k dispozici")
        if fmt == "pdf_ocr":
            msg.setText("Pro vytvoření PDF s rozpoznaným textem je nejprve potřeba provést OCR.")
        elif fmt == "txt":
            msg.setText("Textový obsah zatím není k dispozici, protože nebylo provedeno OCR.")
        else:
            msg.setText("Obsah pro Word dokument zatím není k dispozici, protože nebylo provedeno OCR.")
        msg.setInformativeText(f"Požadovaný formát: {fmt_human}. Co chcete udělat?")
        btn_run = msg.addButton("Spustit OCR", QMessageBox.ButtonRole.YesRole)
        btn_run.setAccessibleName("Spustit OCR a poté uložit")
        btn_run.setAccessibleDescription("Explicitně spustí OCR a po dokončení automaticky uloží dokument.")
        if fmt == "pdf_ocr":
            btn_without = msg.addButton("Uložit PDF bez OCR", QMessageBox.ButtonRole.ActionRole)
            btn_without.setAccessibleName("Uložit PDF bez OCR – pouze obrázky")
        else:
            btn_without = None
        btn_cancel = msg.addButton("Zrušit", QMessageBox.ButtonRole.NoRole)
        btn_cancel.setAccessibleName("Zrušit ukládání")
        msg.setDefaultButton(btn_run)
        # Přístupnost: fokus na Spustit OCR
        msg.exec()
        clicked = msg.clickedButton()
        if clicked == btn_run:
            return "run_ocr"
        if btn_without is not None and clicked == btn_without:
            return "save_without"
        return "cancel"

    def _ask_partial_ocr(self, missing: int, total: int, fmt: str) -> str:
        speak("Některé stránky nemají OCR.")
        msg = QMessageBox(self)
        msg.setWindowTitle("Chybějící OCR")
        msg.setText(f"Některé stránky ještě nemají OCR ({total-missing} z {total} má OCR, chybí {missing}). Chcete nejprve spustit OCR pro chybějící stránky?")
        fmt_human = {"txt": "TXT", "docx": "DOCX", "pdf_ocr": "PDF s OCR"}[fmt]
        msg.setInformativeText(f"Požadovaný formát: {fmt_human}. Chybějící stránky lze rozpoznat doplňkovým OCR.")
        btn_run = msg.addButton("Spustit OCR", QMessageBox.ButtonRole.YesRole)
        btn_run.setAccessibleName("Spustit OCR pro chybějící stránky a poté uložit")
        if fmt == "pdf_ocr":
            btn_without = msg.addButton("Uložit bez OCR", QMessageBox.ButtonRole.ActionRole)
            btn_without.setAccessibleName("Uložit všechny stránky bez OCR")
        else:
            btn_without = None
        btn_cancel = msg.addButton("Zrušit", QMessageBox.ButtonRole.NoRole)
        btn_cancel.setAccessibleName("Zrušit ukládání")
        msg.setDefaultButton(btn_run)
        msg.exec()
        clicked = msg.clickedButton()
        if clicked == btn_run:
            return "run_ocr"
        if btn_without is not None and clicked == btn_without:
            return "save_without"
        return "cancel"

    def _run_ocr_for_pending_save(self, fmt: str) -> None:
        """Spustí OCR explicitně a po úspěchu automaticky uloží pending formát."""
        # zachovej fmt v self._pending_save_format (už nastaveno)
        total_before = len(self.scanned_images)
        try:
            if self._is_ocr_empty():
                # full OCR
                self._ocr_total_pages = total_before
                self._ocr_last_announced_page = 0
                speak(f"Spouštím OCR pro uložení {fmt}, celkem {total_before} stránek.")
                self._run_ocr_flow(interactive=False)
                # _run_ocr_flow v non-interactive neukazuje preview, ale nastaví _last_text
                # pokud našlo text, pokračuj k uložení
            else:
                # incremental – doplnit chybějící
                missing_start = len(self.last_ocr_results)
                self._run_ocr_incremental(missing_start)
                # _run_ocr_incremental už vrací pokud pending, bez preview
        except Exception as e:
            QMessageBox.critical(self, "Chyba OCR", f"OCR selhalo:\n{e}")
            speak("OCR selhalo.")
            self._pending_save_format = None
            return
        # Ověř že OCR nyní kompletní
        if self._is_ocr_empty() or not self._has_ocr_text():
            QMessageBox.warning(self, "OCR bez výsledku", "OCR bylo dokončeno, ale nebyl rozpoznán žádný text. Dokument nebude uložen.")
            speak("OCR neprodukovalo text.")
            self._pending_save_format = None
            return
        # Automaticky pokračovat v původně zvoleném ukládání
        try:
            self._do_save_ocr_format(fmt)
        finally:
            self._pending_save_format = None

    def _do_save_ocr_format(self, fmt: str) -> None:
        """Uloží OCR formát – voláno pouze když OCR existuje, nikdy nespouští OCR."""
        # vyber filtr a příponu
        filters = {"txt": "Textový dokument (*.txt)", "docx": "Word dokument (*.docx)", "pdf_ocr": "PDF s rozpoznaným textem (*.pdf)"}
        suffixes = {"txt": ".txt", "docx": ".docx", "pdf_ocr": ".pdf"}
        chosen_filter = filters[fmt]
        suffix = suffixes[fmt]
        default_dir = QStandardPaths.writableLocation(QStandardPaths.StandardLocation.DocumentsLocation)
        if not default_dir:
            default_dir = os.path.expanduser("~")
        name_map = {"txt": "naskenovany_dokument.txt", "docx": "naskenovany_dokument.docx", "pdf_ocr": "naskenovany_dokument.pdf"}
        default_path = os.path.join(default_dir, name_map[fmt])
        speak("Vyberte umístění pro uložení souboru.")
        path, _ = QFileDialog.getSaveFileName(self, "Uložit dokument", default_path, chosen_filter)
        if not path:
            self.page_list.setFocus()
            return
        if not path.lower().endswith(suffix):
            base, ext = os.path.splitext(path)
            if ext.lower() not in (".txt", ".pdf", ".docx"):
                path = path + suffix
            elif ext.lower() != suffix:
                path = base + suffix
        if os.path.exists(path):
            if not self._confirm_overwrite(path):
                speak("Ukládání zrušeno.")
                self.page_list.setFocus()
                return
        try:
            if fmt == "pdf_ocr":
                self._save_pdf_with_ocr(path)
            elif fmt == "docx":
                self._save_docx_pages(path)
            else:
                # pro TXT použij _last_text (již obsahuje headery)
                self._save_txt_pages(path, self._last_text)
        except Exception as e:
            QMessageBox.critical(self, "Chyba při ukládání", f"Nepodařilo se uložit soubor:\n{e}")
            speak("Chyba při ukládání souboru.")
            self.page_list.setFocus()

    def _confirm_overwrite(self, path: str) -> bool:
        msg = QMessageBox(self)
        msg.setWindowTitle("Soubor již existuje")
        msg.setText(f"Soubor již existuje:\n{path}")
        msg.setInformativeText("Chcete jej přepsat?")
        btn_over = msg.addButton("Přepsat", QMessageBox.ButtonRole.YesRole)
        btn_over.setAccessibleName("Ano, přepsat soubor")
        btn_cancel = msg.addButton("Zrušit", QMessageBox.ButtonRole.NoRole)
        btn_cancel.setAccessibleName("Zrušit, neukládat")
        msg.setDefaultButton(btn_cancel)
        msg.exec()
        return msg.clickedButton() == btn_over

    def read_last_text(self) -> None:
        # Při zapnuté opravě čte processed_text (už v _last_text), jinak originál
        text_to_read = self._last_text
        if self._last_text:
            speak(text_to_read)
        else:
            speak("Není žádný text k přečtení.")

    # ---------- Export ----------
    def export_images(self) -> None:
        if not self.scanned_images:
            QMessageBox.warning(self, "Chyba", "Nejdříve naskenujte stránky.")
            return

        formats = [("JPEG", ".jpg"), ("PNG", ".png"), ("TIFF", ".tif")]
        fmt_items = [f"{name} (*{ext})" for name, ext in formats]
        fmt_combo = QComboBox()
        fmt_combo.addItems(fmt_items)
        fmt_combo.setAccessibleName("Formát obrázků")

        msg = QMessageBox(self)
        msg.setWindowTitle("Export obrázků")
        msg.setText("Zvolte formát obrázků:")
        msg.setInformativeText("Poté budete vyzváni k výběru cílové složky.")
        msg.layout().addWidget(fmt_combo)
        msg.addButton(QMessageBox.StandardButton.Ok)
        msg.addButton(QMessageBox.StandardButton.Cancel)

        if msg.exec() != QMessageBox.StandardButton.Ok:
            return

        dir_path = QFileDialog.getExistingDirectory(self, "Vyberte cílovou složku")
        if not dir_path:
            return

        sel_idx = fmt_combo.currentIndex()
        _, ext = formats[sel_idx]
        pil_fmt = formats[sel_idx][0]

        for i, img in enumerate(self.scanned_images):
            fname = os.path.join(dir_path, f"stranka_{i+1:03d}{ext}")
            if pil_fmt == "JPEG":
                img = img.convert("RGB")
                img.save(fname, pil_fmt, quality=95)
            else:
                img.save(fname, pil_fmt)

        QMessageBox.information(
            self, "Hotovo",
            f"{len(self.scanned_images)} obrázků uloženo do {dir_path}"
        )
        speak("Obrázky exportovány.")

    # ---------- Save helpers ----------
    def _get_docx_pages(self) -> list[str]:
        """Vrátí list textů per-page pro DOCX podle skutečných stránek.

        Priorita:
        1) _last_edited_pages pokud délka == scanned_images
        2) rekonstrukce z last_ocr_results (display_text)
        3) fallback rozdělení _last_text
        """
        total = len(self.scanned_images)
        # 1) upravené stránky pokud sedí počet
        if self._last_edited_pages is not None and len(self._last_edited_pages) == total:
            return list(self._last_edited_pages)
        # Zkusit re-split _last_text pokud obsahuje headery a sedí
        if self._last_text:
            pages = _split_text_by_page_headers(self._last_text)
            if pages is not None and len(pages) == total:
                return pages
        # 2) rekonstrukce z OCR výsledků
        if self.last_ocr_results and len(self.last_ocr_results) == total:
            out: list[str] = []
            for page in self.last_ocr_results:
                if not page:
                    out.append("")
                else:
                    # Spoj display_text jednotlivých OcrResult s novým řádkem
                    # (zachová odřádkování, ne přidává umělé \n\n mezi každým slovem)
                    texts = [r.display_text for r in page]
                    # Pokud je více výsledků, spoj je novým řádkem – DOCX pak splitne \n\n na odstavce
                    out.append("\n".join(texts) if len(texts) > 1 else (texts[0] if texts else ""))
            return out
        # 3) fallback – rozděl _last_text i když nesedí, nebo vrať jako jednu stránku
        if self._last_text:
            pages = _split_text_by_page_headers(self._last_text)
            if pages is not None:
                # Doplň/zkrát na total
                if len(pages) < total:
                    pages = pages + [""] * (total - len(pages))
                return pages[:total]
            return [self._last_text]
        return [""] * total if total else []

    def _save_txt_pages(self, path: str, edited_text: str) -> None:
        """Uloží TXT s oddělovači --- Stránka N --- v pořadí 1..N, čitelné pro NVDA."""
        total = len(self.scanned_images)
        # Pokud edited_text již obsahuje headery a počet sedí nebo je alespoň 1, ulož přímo
        pages = _split_text_by_page_headers(edited_text)
        if pages is not None and len(pages) == total:
            # Obsah již má správné headery – ulož edited_text přímo (zachová přesnou editaci)
            # Ale normalizuj aby každý header byl přesně "--- Stránka N ---"
            # Pro jednoduchost ulož přímo edited_text pokud obsahuje headery
            content = edited_text
            # Zajisti, že soubor končí newline
            if not content.endswith("\n"):
                content += "\n"
        elif total > 0:
            # Generuj per-page z edited_text per-page pokud sedí, jinak fallback na edited split nebo OCR
            if pages is not None and len(pages) == total:
                content = ""
                for i, p in enumerate(pages):
                    content += f"--- Stránka {i+1} ---\n{p}\n\n"
            elif self._last_edited_pages is not None and len(self._last_edited_pages) == total:
                content = ""
                for i, p in enumerate(self._last_edited_pages):
                    content += f"--- Stránka {i+1} ---\n{p}\n\n"
            else:
                # Pokud edited_text nemá headery ale máme total, zkusit rozdělitEdited nebo použít celý edited_text jako stránku 1 + prázdné?
                # Nejbezpečnější: pokud edited_text neobsahuje headery a je jen jeden dokument,
                # ulož ho jako souvislý text s headery podle skutečných stránek.
                # Pokud edited_text obsahuje více odstavců ale bez headerů, nelze bezpečně rekonstruovat per-page,
                # takže pokud total==1 ulož přímo, jinak vygeneruj z OCR nebo z edited_text jako celek pro stránku 1.
                if total == 1:
                    # Pokud jedna stránka, bez headeru je ok, ale pro konzistenci přidej header
                    # Pokud uživatel explicitně smazal header, respektuj jeho editaci – ulož bez headeru
                    # Detekce: pokud edited_text nemá header a total==1, ulož edited_text přímo
                    content = edited_text
                else:
                    # Více stránek ale text bez headerů – pokus se generovat z OCR per-page pokud dostupné
                    if self.last_ocr_results and len(self.last_ocr_results) == total:
                        # Pokud je edited_text výrazně odlišný od OCR, může být uživatelská editace
                        # – v tom případě bez spolehlivého per-page rozdělení je bezpečnější uložit
                        # edited_text jako celek bez umělého dělení.
                        # Preferuj edited_text jako celek pro TXT čitelnost, ale zachovej informaci o stránkách:
                        # Zkusíme: pokud edited_text obsahuje "\n\n" ale ne headery, uložíme ho s headery pouze pro první stránku?
                        # Ne – specifikace říká TXT může obsahovat souvislý text všech stránek s oddělovači.
                        # Pokud nemáme per-page, použij OCR per-page jako zdroj a ignoruj edited? To by zahodilo editaci.
                        # Proto: ulož edited_text jako souvislý text s fallback generováním oddělovačů jen pokud edited_text je per-page.
                        # Aktuálně: edited_text bez headerů pro více stran → ulož edited_text přímo (bez umělého dělení) + přidej info?
                        # Rozhodnutí: ulož edited_text přímo, protože umělé dělení by bylo nebezpečné.
                        # Ale pro konzistenci a NVDA čitelnost je lepší mít headery. Pokud edited_text nemá headery,
                        # vytvoř per-page z OCR a nepoužívej edited – dokument by ztratil editaci. To je horší.
                        # Kompromis: pokud edited_text bez headerů a total>1, ulož s headery kde stránka 1 = edited_text, ostatní prázdné?
                        # Ne, to je matoucí. Nejmenší překvapení: ulož edited_text přímo.
                        content = edited_text
                    else:
                        content = edited_text
        else:
            content = edited_text

        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        self._show_save_success(path, "Textový dokument (.txt)", "Textový dokument")

    def _save_pdf_without_ocr(self, path: str) -> None:
        """Uloží stránky jako obrázky. OCR nebude proveden. Použije přímo self.scanned_images."""
        doc = fitz.open()
        dpi = self.dpi_combo.currentData()
        if dpi is None:
            dpi = 300
        for img in self.scanned_images:
            page = doc.new_page(width=img.width / dpi * 72, height=img.height / dpi * 72)
            tmp_path = None
            try:
                with tempfile.NamedTemporaryFile(delete=False, suffix=".jpg") as tmp:
                    tmp_path = tmp.name
                    img.save(tmp_path, "JPEG", quality=95)
                page.insert_image(page.rect, filename=tmp_path)
            finally:
                if tmp_path and os.path.exists(tmp_path):
                    try:
                        os.remove(tmp_path)
                    except OSError:
                        pass
        doc.save(path)
        self._show_save_success(path, "PDF – naskenované stránky (.pdf)", "PDF bez OCR")

    def _save_pdf_with_ocr(self, path: str) -> None:
        """Zachová vzhled naskenovaných stránek a přidá textovou vrstvu z OCR. Použije self.last_ocr_results, nespouští OCR."""
        doc = fitz.open()
        dpi = self.dpi_combo.currentData()
        if dpi is None:
            dpi = 300
        for i, img in enumerate(self.scanned_images):
            page = doc.new_page(width=img.width / dpi * 72, height=img.height / dpi * 72)
            tmp_path = None
            try:
                with tempfile.NamedTemporaryFile(delete=False, suffix=".jpg") as tmp:
                    tmp_path = tmp.name
                    img.save(tmp_path, "JPEG", quality=95)
                page.insert_image(page.rect, filename=tmp_path)
                if i < len(self.last_ocr_results):
                    for item in self.last_ocr_results[i]:
                        if item.bbox is not None:
                            bbox = item.bbox
                            x0, y0 = bbox[0][0], bbox[0][1]
                            x1, y1 = bbox[2][0], bbox[2][1]
                            rect = fitz.Rect(x0 / dpi * 72, y0 / dpi * 72, x1 / dpi * 72, y1 / dpi * 72)
                            txt = item.display_text if hasattr(item, "display_text") else item.text
                            page.insert_textbox(rect, txt, fontsize=0, fill_opacity=0)
            finally:
                if tmp_path and os.path.exists(tmp_path):
                    try:
                        os.remove(tmp_path)
                    except OSError:
                        pass
        doc.save(path)
        self._show_save_success(path, "PDF s rozpoznaným textem (.pdf)", "PDF s OCR")

    def _save_pdf(self, path: str) -> None:
        """Legacy alias – pro kompatibilitu volá _save_pdf_with_ocr pokud existuje OCR, jinak bez."""
        if self._is_ocr_complete():
            self._save_pdf_with_ocr(path)
        else:
            # zachovej původní chování – vlož OCR vrstvu pokud existuje částečně
            # ale nový tok by měl volat explicitně jednu z variant
            if self.last_ocr_results:
                self._save_pdf_with_ocr(path)
            else:
                self._save_pdf_without_ocr(path)

    def _save_docx(self, text: str, path: str) -> None:
        """Legacy wrapper – zachován pro kompatibilitu (makra). Nově volá per-page logiku."""
        # Aktualizuj _last_text a _last_edited_pages pro per-page logiku
        self._last_text = text
        pages = _split_text_by_page_headers(text)
        if pages is not None:
            self._last_edited_pages = pages
        self._save_docx_pages(path)

    def _save_docx_pages(self, path: str) -> None:
        """Uloží DOCX rozdělený podle skutečných naskenovaných stránek (page break pouze mezi stránkami)."""
        doc = Document()
        pages = self._get_docx_pages()
        total = len(pages)
        for idx, page_text in enumerate(pages):
            # Rozděl na odstavce podle dvojitého odřádkování, ale bez přidávání page break mezi odstavci
            if page_text is None:
                page_text = ""
            # Zachovej prázdné stránky
            if not page_text.strip():
                doc.add_paragraph("")
            else:
                paragraphs = page_text.split("\n\n")
                for para in paragraphs:
                    # Prázdné odstavce zachovej jako prázdný paragraph
                    doc.add_paragraph(para)
            if idx < total - 1:
                doc.add_page_break()
        doc.save(path)
        self._show_save_success(path, "Word dokument (.docx)", "Word dokument")

    def _show_save_success(self, path: str, format_human: str, format_short: str) -> None:
        """Společný úspěšný dialog – přístupný, s cestou v textu, krátkou hlasovou hláškou."""
        QMessageBox.information(
            self, "Hotovo",
            f"Dokument byl úspěšně uložen.\nFormát: {format_human}\nCesta: {path}"
        )
        # Hlasově krátce, bez celé cesty
        speak(f"Dokument byl úspěšně uložen. Dokument byl uložen jako {format_short}.")
        self.btn_scan.setFocus()

    # ---------- Macros ----------
    def _rebuild_macro_buttons(self) -> None:
        while self._macro_buttons_layout.count():
            item = self._macro_buttons_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()
        for macro in self.macro_manager.macros:
            container = QWidget()
            row = QHBoxLayout(container)
            row.setContentsMargins(0, 0, 0, 0)
            label = macro.name
            if macro.shortcut:
                label += f" ({macro.shortcut})"
            btn_run = QPushButton(label)
            btn_run.setAccessibleName(f"Spustit makro: {macro.name}")
            btn_run.clicked.connect(lambda checked, m=macro: self._run_macro(m))
            row.addWidget(btn_run)

            btn_edit = QPushButton("Upravit")
            btn_edit.setAccessibleName(f"Upravit makro: {macro.name}")
            btn_edit.clicked.connect(lambda checked, m=macro: self._open_macro_editor(m))
            row.addWidget(btn_edit)

            btn_rename = QPushButton("Přejmenovat")
            btn_rename.setAccessibleName(f"Přejmenovat makro: {macro.name}")
            btn_rename.clicked.connect(lambda checked, m=macro: self._rename_macro(m))
            row.addWidget(btn_rename)

            btn_delete = QPushButton("Smazat")
            btn_delete.setAccessibleName(f"Smazat makro: {macro.name}")
            btn_delete.clicked.connect(lambda checked, m=macro: self._delete_macro(m))
            row.addWidget(btn_delete)

            self._macro_buttons_layout.addWidget(container)

    def _open_macro_editor(self, macro: Optional[Macro] = None) -> None:
        dialog = MacroEditorDialog(self, self.macro_manager, macro)
        dialog.exec()

    def _run_macro(self, macro: Macro) -> None:
        self.speak(f"Spouštím makro: {macro.name}")
        self._macro_runner.run(macro)
        self.speak("Makro dokončeno.")

    def _rename_macro(self, macro: Macro) -> None:
        new_name, ok = QInputDialog.getText(
            self, "Přejmenovat makro", "Nový název makra:",
            text=macro.name,
        )
        if ok and new_name.strip():
            self.macro_manager.rename(macro, new_name.strip())
            self.macro_manager.load_all()
            self._rebuild_macro_buttons()
            self._setup_macro_shortcuts()
            self.speak(f"Makro přejmenováno na {new_name.strip()}")

    def _delete_macro(self, macro: Macro) -> None:
        msg = QMessageBox(self)
        msg.setWindowTitle("Smazat makro")
        msg.setText(f"Opravdu chcete smazat makro {macro.name}?")
        msg.setInformativeText("Tuto akci nelze vrátit.")
        btn_yes = msg.addButton("Ano", QMessageBox.ButtonRole.YesRole)
        btn_yes.setAccessibleName(f"Ano, smazat makro {macro.name}")
        btn_no = msg.addButton("Ne", QMessageBox.ButtonRole.NoRole)
        btn_no.setAccessibleName("Ne, ponechat makro")
        msg.exec()
        if msg.clickedButton() == btn_yes:
            self.macro_manager.delete(macro)
            self._rebuild_macro_buttons()
            self._setup_macro_shortcuts()
            self.speak(f"Makro {macro.name} smazáno.")


if __name__ == "__main__":
    app = QApplication(sys.argv)
    font = app.font()
    font.setPointSize(11)
    app.setFont(font)
    win = ScanApp()
    win.show()
    sys.exit(app.exec())
